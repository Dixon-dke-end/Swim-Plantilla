"""Generate the income Excel template and the sample files used by /contabilidad.

Outputs (under public/plantillas/):
  - plantilla-ingresos-marea.xlsx   official template (instructions, dropdowns, formulas)
  - ejemplos/ventas-marea-2026.xlsx  six months of sales in the official format
  - ejemplos/cuadre-caja-pos.xlsx    point-of-sale export with title rows and other headers
  - ejemplos/extracto-bancario.csv   bank statement (semicolon CSV, Colombian number format)

Run: python -I scripts/generate_income_templates.py
"""

from __future__ import annotations

import csv
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "public" / "plantillas"
SAMPLES = OUT / "ejemplos"

TEAL = "0D4A4C"
SAND = "F3EDE4"
INK = "13201F"
CORAL = "E3643F"
COP = '"$"#,##0'
DATE = "dd/mm/yyyy"
IVA_RATE = 0.19

PRODUCTS = [
    ("Bikini Alba", "Bikinis", 189_000),
    ("Enterizo Ola", "Enterizos", 289_000),
    ("Bikini Lunares", "Bikinis", 179_000),
    ("Short Costa", "Hombre", 169_000),
    ("Enterizo Lago", "Enterizos", 269_000),
    ("Enterizo Sol", "Enterizos", 319_000),
    ("Bikini Palma", "Bikinis", 209_000),
    ("Bikini Roca", "Bikinis", 159_000),
    ("Enterizo Jardín", "Resort", 349_000),
    ("Short Piscina", "Hombre", 175_000),
    ("Short Girasol", "Hombre", 185_000),
    ("Enterizo Oasis", "Resort", 389_000),
]
CHANNELS = ["Tienda online", "Tienda física", "Instagram", "Mayorista"]
CATEGORIES = ["Bikinis", "Enterizos", "Hombre", "Resort", "Accesorios"]
PAYMENTS = ["Tarjeta", "PSE", "Nequi", "Efectivo", "Transferencia"]
STATUSES = ["Pagado", "Pendiente", "Anulado"]
FIRST_NAMES = [
    "Valentina", "Camila", "Mariana", "Daniela", "Sofía", "Laura", "Isabella",
    "Andrés", "Santiago", "Juliana", "Natalia", "Carolina", "Felipe", "Paula",
]
LAST_INITIALS = "ABCDEGHJLMPRSTV"
WHOLESALE = ["Boutique Arena SAS", "Playa Store", "Moda Caribe Ltda.", "Hotel Brisa Marina"]

HEADERS = [
    ("Fecha", 13), ("N° factura", 13), ("Cliente", 24), ("Canal", 15),
    ("Categoría", 13), ("Producto", 20), ("Cantidad", 10),
    ("Precio unitario (IVA incl.)", 16), ("Descuento", 13), ("Total", 15),
    ("Base gravable", 15), ("IVA 19%", 13), ("Medio de pago", 15),
    ("Estado", 12), ("Fecha de pago", 14), ("Notas", 30),
]

thin = Side(style="thin", color="D9D1C4")
header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill("solid", fgColor=TEAL)


def style_header(ws, row: int, columns: list[tuple[str, int]]) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=index, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[row].height = 32


def apply_income_formats(ws, first: int, last: int) -> None:
    for row in range(first, last + 1):
        ws.cell(row=row, column=1).number_format = DATE
        ws.cell(row=row, column=15).number_format = DATE
        for col in (8, 9, 10, 11, 12):
            ws.cell(row=row, column=col).number_format = COP


def add_lists_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Listas")
    lists = [("Canal", CHANNELS), ("Categoría", CATEGORIES), ("Medio de pago", PAYMENTS), ("Estado", STATUSES)]
    for col, (title, values) in enumerate(lists, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        ws.column_dimensions[get_column_letter(col)].width = 18
        for row, value in enumerate(values, start=2):
            ws.cell(row=row, column=col, value=value)


def add_validations(ws, last_row: int) -> None:
    ranges = {"D": "$A$2:$A$5", "E": "$B$2:$B$6", "M": "$C$2:$C$6", "N": "$D$2:$D$4"}
    for column, source in ranges.items():
        validation = DataValidation(type="list", formula1=f"=Listas!{source}", allow_blank=True)
        validation.error = "Elige un valor de la lista."
        validation.errorTitle = "Valor no válido"
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}{last_row}")


def add_instructions_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Instrucciones"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 80

    ws["B2"] = "MAREA · Plantilla de ingresos"
    ws["B2"].font = Font(name="Georgia", size=22, color=TEAL, bold=True)
    ws["B3"] = "Registra cada venta en la hoja «Ingresos» y súbela en /contabilidad."
    ws["B3"].font = Font(size=12, color=INK)

    rows = [
        ("Columna", "Qué escribir"),
        ("Fecha", "Fecha de la venta o factura (dd/mm/aaaa). Obligatoria."),
        ("N° factura", "Consecutivo de la factura o ticket. Sirve para detectar duplicados."),
        ("Cliente", "Nombre del cliente o razón social."),
        ("Canal", "Tienda online, Tienda física, Instagram o Mayorista (lista desplegable)."),
        ("Categoría", "Bikinis, Enterizos, Hombre, Resort o Accesorios."),
        ("Producto", "Nombre del producto vendido."),
        ("Cantidad", "Unidades vendidas."),
        ("Precio unitario", "Precio de venta al público con IVA incluido."),
        ("Descuento", "Valor total descontado en la línea (opcional)."),
        ("Total / Base / IVA", "Se calculan solos. No los escribas a mano."),
        ("Medio de pago", "Tarjeta, PSE, Nequi, Efectivo o Transferencia."),
        ("Estado", "Pagado, Pendiente o Anulado. Las anuladas no suman."),
        ("Fecha de pago", "Cuándo entró el dinero. Se usa en contabilidad por caja."),
        ("Notas", "Cualquier observación."),
    ]
    for offset, (left, right) in enumerate(rows):
        row = 5 + offset
        ws.cell(row=row, column=2, value=left)
        ws.cell(row=row, column=3, value=right)
        if offset == 0:
            for col in (2, 3):
                cell = ws.cell(row=row, column=col)
                cell.font = header_font
                cell.fill = header_fill
        else:
            ws.cell(row=row, column=2).font = Font(bold=True, color=INK)
            for col in (2, 3):
                ws.cell(row=row, column=col).border = Border(bottom=thin)

    note_row = 5 + len(rows) + 1
    ws.cell(row=note_row, column=2, value="¿Tu contador usa otro formato?")
    ws.cell(row=note_row, column=2).font = Font(bold=True, color=CORAL, size=12)
    ws.cell(
        row=note_row + 1,
        column=2,
        value=(
            "No hace falta copiar los datos a esta plantilla: /contabilidad reconoce las columnas de "
            "extractos bancarios, cuadres de caja y libros de ventas, y guarda el mapeo como perfil."
        ),
    )
    ws.merge_cells(start_row=note_row + 1, start_column=2, end_row=note_row + 1, end_column=3)
    ws.cell(row=note_row + 1, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_row + 1].height = 36


def build_template() -> None:
    wb = Workbook()
    add_instructions_sheet(wb)
    ws = wb.create_sheet("Ingresos")
    style_header(ws, 1, HEADERS)
    ws.freeze_panes = "A2"

    last_row = 1001
    examples = [
        (date(2026, 10, 1), "FV-0001", "Cliente de ejemplo", "Tienda online", "Bikinis", "Bikini Alba", 1, 189_000, 0, "Tarjeta", "Pagado", date(2026, 10, 1), "Fila de ejemplo: puedes borrarla"),
        (date(2026, 10, 2), "FV-0002", "Boutique Arena SAS", "Mayorista", "Enterizos", "Enterizo Ola", 6, 289_000, 173_400, "Transferencia", "Pendiente", None, "Pedido mayorista con 10 % de descuento"),
    ]
    for offset, example in enumerate(examples):
        row = 2 + offset
        when, invoice, client, channel, category, product, qty, price, discount, payment, status, paid, notes = example
        values = {1: when, 2: invoice, 3: client, 4: channel, 5: category, 6: product, 7: qty, 8: price, 9: discount, 13: payment, 14: status, 15: paid, 16: notes}
        for col, value in values.items():
            ws.cell(row=row, column=col, value=value)

    for row in range(2, last_row + 1):
        ws.cell(row=row, column=10, value=f'=IF(OR(G{row}="",H{row}=""),"",G{row}*H{row}-N(I{row}))')
        ws.cell(row=row, column=11, value=f'=IF(J{row}="","",ROUND(J{row}/(1+{IVA_RATE}),0))')
        ws.cell(row=row, column=12, value=f'=IF(J{row}="","",J{row}-K{row})')
        for col in (10, 11, 12):
            ws.cell(row=row, column=col).font = Font(color="5B6B69")
    apply_income_formats(ws, 2, last_row)

    table = Table(displayName="Ingresos", ref=f"A1:P{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)

    add_lists_sheet(wb)
    add_validations(ws, last_row)
    wb.active = 1
    wb.save(OUT / "plantilla-ingresos-marea.xlsx")


def monthly_volume(month: int) -> int:
    # Mid-year school holidays (June/July) and the October launch drive peaks.
    return {5: 34, 6: 52, 7: 58, 8: 38, 9: 36, 10: 44}[month]


def client_name(rng: random.Random) -> str:
    return f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_INITIALS)}."


def generate_sales(rng: random.Random) -> list[dict]:
    sales: list[dict] = []
    invoice = 1040
    for month in range(5, 11):
        days = (date(2026, month + 1, 1) - timedelta(days=1)).day if month < 12 else 31
        if month == 10:
            days = 17
        for _ in range(monthly_volume(month)):
            invoice += 1
            when = date(2026, month, rng.randint(1, days))
            product, category, price = rng.choice(PRODUCTS)
            channel = rng.choices(CHANNELS, weights=[46, 26, 20, 8])[0]
            quantity = rng.randint(4, 12) if channel == "Mayorista" else rng.choices([1, 2, 3], weights=[78, 18, 4])[0]
            discount = round(quantity * price * 0.1) if channel == "Mayorista" else rng.choice([0, 0, 0, 0, 20_000])
            total = quantity * price - discount
            base = round(total / (1 + IVA_RATE))
            payment = "Transferencia" if channel == "Mayorista" else rng.choices(PAYMENTS, weights=[38, 22, 20, 14, 6])[0]
            if channel == "Tienda física" and payment == "PSE":
                payment = "Efectivo"
            status = rng.choices(STATUSES, weights=[90, 7, 3])[0]
            if channel == "Mayorista" and rng.random() < 0.4:
                status = "Pendiente"
            paid = None
            if status == "Pagado":
                lag = rng.randint(15, 30) if channel == "Mayorista" else rng.choice([0, 0, 0, 1])
                paid = min(when + timedelta(days=lag), date(2026, 10, 17))
            sales.append({
                "date": when, "invoice": f"FV-{invoice:04d}",
                "client": rng.choice(WHOLESALE) if channel == "Mayorista" else client_name(rng),
                "channel": channel, "category": category, "product": product,
                "quantity": quantity, "price": price, "discount": discount,
                "total": total, "base": base, "iva": total - base,
                "payment": payment, "status": status, "paid": paid,
                "notes": "Pedido mayorista" if channel == "Mayorista" else "",
            })
    sales.sort(key=lambda sale: (sale["date"], sale["invoice"]))
    return sales


def build_sales_sample(sales: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ingresos"
    style_header(ws, 1, HEADERS)
    ws.freeze_panes = "A2"
    for row, sale in enumerate(sales, start=2):
        values = [
            sale["date"], sale["invoice"], sale["client"], sale["channel"], sale["category"],
            sale["product"], sale["quantity"], sale["price"], sale["discount"], sale["total"],
            sale["base"], sale["iva"], sale["payment"], sale["status"], sale["paid"], sale["notes"],
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
    last = len(sales) + 1

    # Rows with real-world problems so the importer can show its validation report.
    problems = [
        [date(2026, 9, 12), "FV-9001", "Laura M.", "Instagram", "Bikinis", "Bikini Palma", 1, None, 0, None, None, None, "Nequi", "Pagado", date(2026, 9, 12), "Falta el precio"],
        ["31/02/2026", "FV-9002", "Paula R.", "Tienda online", "Enterizos", "Enterizo Sol", 1, 319_000, 0, 319_000, 268_067, 50_933, "Tarjeta", "Pagado", None, "Fecha imposible"],
        [sales[10]["date"], sales[10]["invoice"], sales[10]["client"], sales[10]["channel"], sales[10]["category"], sales[10]["product"], sales[10]["quantity"], sales[10]["price"], sales[10]["discount"], sales[10]["total"], sales[10]["base"], sales[10]["iva"], sales[10]["payment"], sales[10]["status"], sales[10]["paid"], "Factura duplicada"],
    ]
    for offset, values in enumerate(problems, start=1):
        for col, value in enumerate(values, start=1):
            ws.cell(row=last + offset, column=col, value=value)
    last += len(problems)

    apply_income_formats(ws, 2, last)
    table = Table(displayName="Ventas2026", ref=f"A1:P{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    add_lists_sheet(wb)
    wb.save(SAMPLES / "ventas-marea-2026.xlsx")


def build_pos_sample(rng: random.Random) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cuadre"
    ws["A1"] = "CUADRE DE CAJA · PUNTO DE VENTA CC SANTAFÉ"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Generado por el software POS · Periodo: agosto - octubre 2026"
    ws["A2"].font = Font(italic=True, color="666666")
    columns = [("Fecha/Hora", 19), ("Ticket", 10), ("Vendedor", 14), ("Forma de pago", 15), ("Valor neto", 14), ("Impuesto", 13), ("Total venta", 14)]
    style_header(ws, 4, columns)
    sellers = ["Marcela", "Jhon", "Diana"]
    row = 5
    ticket = 5120
    start = datetime(2026, 8, 1, 10, 0)
    while start < datetime(2026, 10, 17):
        for _ in range(rng.randint(0, 3)):
            ticket += 1
            product = rng.choice(PRODUCTS)
            total = product[2] * rng.choices([1, 2], weights=[85, 15])[0]
            net = round(total / (1 + IVA_RATE))
            moment = start + timedelta(hours=rng.randint(0, 9), minutes=rng.randint(0, 59))
            values = [moment, f"T-{ticket}", rng.choice(sellers), rng.choice(["Efectivo", "Tarjeta débito", "Tarjeta crédito", "Nequi"]), net, total - net, total]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                if col == 1:
                    cell.number_format = "dd/mm/yyyy hh:mm"
                if col >= 5:
                    cell.number_format = COP
            row += 1
        start += timedelta(days=1)
    ws.cell(row=row + 1, column=4, value="TOTAL").font = Font(bold=True)
    ws.cell(row=row + 1, column=7, value=f"=SUM(G5:G{row - 1})").number_format = COP
    wb.save(SAMPLES / "cuadre-caja-pos.xlsx")


def colombian_number(value: int) -> str:
    return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def build_bank_sample(rng: random.Random) -> None:
    balance = 18_450_000
    rows = []
    day = date(2026, 9, 1)
    while day <= date(2026, 10, 16):
        for _ in range(rng.randint(1, 3)):
            if rng.random() < 0.72:
                amount = rng.choice(PRODUCTS)[2] * rng.choice([1, 1, 2])
                kind = rng.choice(["PAGO PSE", "ABONO DATÁFONO", "TRANSF NEQUI", "CONSIGNACIÓN"])
                description = f"{kind} {rng.choice(FIRST_NAMES).upper()}"
                debit, credit = 0, amount
            else:
                amount = rng.choice([85_000, 240_000, 1_200_000, 3_450_000])
                description = rng.choice(["PAGO PROVEEDOR TELAS", "PAGO NÓMINA", "CUOTA DE MANEJO", "PAGO ARRIENDO LOCAL"])
                debit, credit = amount, 0
            balance += credit - debit
            rows.append([
                day.isoformat(), description, f"{rng.randint(10_000_000, 99_999_999)}",
                colombian_number(debit) if debit else "", colombian_number(credit) if credit else "",
                colombian_number(balance),
            ])
        day += timedelta(days=1)
    with open(SAMPLES / "extracto-bancario.csv", "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["Fecha transacción", "Descripción", "Referencia", "Débito", "Crédito", "Saldo"])
        writer.writerows(rows)


def main() -> None:
    SAMPLES.mkdir(parents=True, exist_ok=True)
    rng = random.Random(2026)
    build_template()
    sales = generate_sales(rng)
    build_sales_sample(sales)
    build_pos_sample(rng)
    build_bank_sample(rng)
    print(f"{len(sales)} sales generated")


if __name__ == "__main__":
    main()
